import os
from docx import Document

def create_samples():
    os.makedirs("sample_docs", exist_ok=True)
    
    # 1. Create Sample DOCX (Medical)
    doc = Document()
    doc.add_heading('CONFIDENTIAL MEDICAL REPORT', 0)
    doc.add_paragraph('Patient: Jane Smith')
    doc.add_paragraph('SSN: 999-00-1111')
    doc.add_paragraph('Email: smith.jane@provider.net')
    doc.add_paragraph('Symptoms: Severe chest pain and shortness of breath.')
    doc.add_paragraph('Analysis: Patient requires an immediate echocardiogram.')
    doc.save('sample_docs/medical_demo.docx')
    print("Created sample_docs/medical_demo.docx")

    # 2. Create Sample TXT (Financial)
    with open('sample_docs/financial_report.txt', 'w') as f:
        f.write("QUARTERLY REVENUE REPORT\n")
        f.write("Account: 123456789\n")
        f.write("Owner: Michael Scott\n")
        f.write("Total Revenue: $1,200,000\n")
        f.write("Risk: High volatility in paper market.\n")
    print("Created sample_docs/financial_report.txt")

if __name__ == "__main__":
    create_samples()
